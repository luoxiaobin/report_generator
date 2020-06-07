import ExcelReport
import os
import pytest
from openpyxl.reader.excel import load_workbook

def test_excel_report():

    excel_filename = "test.xlsx"
    excel_filename_fullpath = os.path.join(os.path.dirname(os.path.realpath(__file__)), excel_filename)
    os.remove(excel_filename_fullpath) if os.path.exists(excel_filename_fullpath) else None

    column = ['title', 'number', 'description']
    data = [["col1 row2", 12345.67, 0.1234567],
            ["col1 row3", 12345.67, 0.1234567]]          

    ExcelReport.write_to_excel(report_name='test', column=column, data=data,  file_name= excel_filename_fullpath)

    wb = load_workbook(filename = f"{excel_filename_fullpath}") 
    ws = wb["Sheet"]
    
    assert (ws.cell(2, 1).value, ws.cell(3, 1).value)  == ('title', 'col1 row2')
    
    

